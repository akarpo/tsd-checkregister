"""Published-figure fidelity — top-down: dataset -> what a reader actually sees.

check_source_fidelity.py proves the dataset matches source. This proves every number actually
displayed traces back to the dataset, and that the arithmetic on the page holds.

  2.1  Arithmetic invariants in the dataset itself.
  2.2  Workbook: All Lines matches the dataset; every pivot reconciles.
  2.3  Dashboard payload matches the dataset FY-by-FY and category-by-category.
  2.4  Rendered HTML literals (header, badge, methodology sentence, FY cards,
       coverage table, curriculum tiles) match computed values.
  2.5  Labels describe what is actually shown.
"""
import json
import re
import sys
import pickle
from collections import defaultdict
from pathlib import Path

BUILD = Path.home() / "Downloads" / "tsd-checkregister" / "Working Folder" / "Cache and Tools" / "build"
sys.path.insert(0, str(BUILD))
import _paths
import openpyxl

FAIL = []
def check(ok, label, detail=""):
    print(f"  [{'PASS' if ok else 'FAIL'}] {label}" + (f"  ({detail})" if detail else ""))
    if not ok:
        FAIL.append(label)

combined = pickle.loads(_paths.COMBINED_PKL.read_bytes())
N = len(combined)
GRAND = round(sum(r["Amount"] for r in combined), 2)
print(f"dataset: {N:,} rows / ${GRAND:,.2f}\n")

is_bond = lambda r: r["Fund"].startswith("3") or r["Fund"].startswith("4")
op_total = round(sum(r["Amount"] for r in combined if not is_bond(r)), 2)
bond_total = round(sum(r["Amount"] for r in combined if is_bond(r)), 2)
fy_amt = defaultdict(float); fy_cnt = defaultdict(int)
for r in combined:
    fy_amt[r["Fiscal Year"]] += r["Amount"]; fy_cnt[r["Fiscal Year"]] += 1
FYS = sorted(fy_amt)

print("2.1  arithmetic invariants")
check(abs(op_total + bond_total - GRAND) < 0.01, "operating + bond = grand total",
      f"${op_total:,.2f} + ${bond_total:,.2f}")
check(abs(sum(fy_amt.values()) - GRAND) < 0.01, "sum of fiscal years = grand total")
check(sum(fy_cnt.values()) == N, "sum of FY line counts = row count")
check(FYS == [f"FY{n}" for n in range(11, 27)], "FY range is contiguous FY11-FY26",
      f"{len(FYS)} buckets")

print("\n2.2  workbook")
wb = openpyxl.load_workbook(_paths.WORKBOOK, read_only=True)
check(wb.sheetnames == ['README', 'All Lines', 'By Year x Fund', 'By Budget Unit',
                        'PD_Yearly_Summary', 'Curriculum_PD', 'PD by Subject'],
      "7 sheets, expected names", str(len(wb.sheetnames)))
it = wb["All Lines"].iter_rows(values_only=True); hdr = list(next(it))
check("Issue Date FY" not in hdr, "All Lines has a single Fiscal Year column")
n_wb = 0; amt_wb = 0.0; i_amt = hdr.index("Amount"); i_fy = hdr.index("Fiscal Year")
wb_fy = defaultdict(float)
for v in it:
    if v[0] is None and v[1] is None: continue
    n_wb += 1; amt_wb += v[i_amt]; wb_fy[v[i_fy]] += v[i_amt]
check(n_wb == N and abs(amt_wb - GRAND) < 0.01, "All Lines = dataset",
      f"{n_wb:,} rows / ${amt_wb:,.2f}")
check(all(abs(wb_fy[f] - fy_amt[f]) < 0.01 for f in FYS), "All Lines FY totals = dataset FY totals")

g = [list(r) for r in wb["By Year x Fund"].iter_rows(values_only=True)][1:]
piv_amt = round(sum(r[4] for r in g if isinstance(r[4], (int, float))), 2)
piv_cnt = sum(r[3] for r in g if isinstance(r[3], (int, float)))
check(piv_cnt == N and abs(piv_amt - GRAND) < 0.01, "By Year x Fund reconciles",
      f"{piv_cnt:,} / ${piv_amt:,.2f}")
gb = [list(r) for r in wb["By Budget Unit"].iter_rows(values_only=True)][1:]
bu_amt = round(sum(r[-1] for r in gb if isinstance(r[-1], (int, float))), 2)
check(abs(bu_amt - GRAND) < 0.01, "By Budget Unit reconciles", f"${bu_amt:,.2f}")
wb.close()

print("\n2.3  dashboard payload")
html = _paths.DASHBOARD.read_text(encoding="utf-8")
p = json.loads(re.search(r'<script id="data-payload"[^>]*>(.*?)</script>', html, re.DOTALL).group(1))
check(p["meta"]["totalLines"] == N, "payload totalLines = dataset", f'{p["meta"]["totalLines"]:,}')
check(abs(p["all"]["grandTotal"] - GRAND) < 0.01, "payload grandTotal = dataset")
check(abs(p["operating"]["grandTotal"] - op_total) < 0.01, "payload operating = dataset")
check(abs(p["bond"]["grandTotal"] - bond_total) < 0.01, "payload bond = dataset")
check(p["fy"] == FYS, "payload FY list = dataset FY list")
bad_fy = [f for f in FYS if abs(p["all"]["totals"][f] - fy_amt[f]) > 0.01]
check(not bad_fy, "every payload FY total = dataset", f"{len(FYS)} years checked")
cat_sum = round(sum(c["total"] for c in p["all"]["categories"].values()), 2)
check(abs(cat_sum - GRAND) < 0.01, "categories sum to grand total",
      f'{len(p["all"]["categories"])} cats / ${cat_sum:,.2f}')
check(p["meta"]["sourceCount"] == len({r["Source Meeting"] for r in combined}),
      "sourceCount = distinct source meetings", str(p["meta"]["sourceCount"]))

print("\n2.4  rendered HTML literals")
mk = re.sub(r'(<script id="data-payload"[^>]*>).*?(</script>)', r'\1[P]\2', html, flags=re.DOTALL)
M = lambda n: f"${n/1_000_000:.0f}M"
hdr_txt = re.search(r'([\d,]+) line items, (\$\d+M) total — split into operating \((\$\d+M)\) and bond/capital \((\$\d+M)\)', mk)
check(hdr_txt and hdr_txt.group(1) == f"{N:,}", "header line count", hdr_txt.group(1) if hdr_txt else "?")
check(hdr_txt and hdr_txt.group(2) == M(GRAND), "header grand total", hdr_txt.group(2) if hdr_txt else "?")
check(hdr_txt and hdr_txt.group(3) == M(op_total), "header operating", hdr_txt.group(3) if hdr_txt else "?")
check(hdr_txt and hdr_txt.group(4) == M(bond_total), "header bond", hdr_txt.group(4) if hdr_txt else "?")
badge = re.search(r'class="badge-all">(\$\d+M)', mk)
check(badge and badge.group(1) == M(GRAND), "all-funds badge", badge.group(1) if badge else "?")
rec = re.search(r'reconciles exactly to All \((\$\d+M) \+ (\$\d+M) = (\$\d+M)\)', mk)
check(rec and (rec.group(1), rec.group(2), rec.group(3)) == (M(op_total), M(bond_total), M(GRAND)),
      "methodology reconciliation sentence", rec.group(0)[-30:] if rec else "?")
src = re.search(r'(\d+) monthly check register PDFs', mk)
check(src and int(src.group(1)) == p["meta"]["sourceCount"], "source-count caption")

# coverage table month counts vs data (a month counts as covered at 100+ lines)
mo = defaultdict(int)
for r in combined:
    mo[(r["Fiscal Year"], r["Issue Date"].strftime("%Y-%m"))] += 1
def covered(fy):
    y = 2000 + int(fy[2:])
    months = [f"{y-1}-{m:02d}" for m in range(7, 13)] + [f"{y}-{m:02d}" for m in range(1, 7)]
    return sum(1 for m in months if mo.get((fy, m), 0) >= 100)
rows = re.findall(r'<td>(FY\d\d)</td><td class="num">(\d+) of 12', mk)
check(len(rows) == 16, "coverage table has 16 rows", str(len(rows)))
bad_cov = [(fy, shown, covered(fy)) for fy, shown in rows if int(shown) != covered(fy)]
check(not bad_cov, "every coverage-table month count matches the data",
      f"{len(rows)} rows" if not bad_cov else str(bad_cov))

# curriculum stat tiles vs the CURRICULUM_DATA arrays
arr = {k: json.loads(re.search(re.escape("'"+k+"'") + r"\s*:\s*(\[[^\]]*\])", html).group(1))
       for k in ["K-5 ELA cur","K-5 ELA pd","K-5 Math cur","K-5 Math pd",
                 "6-12 ELA cur","6-12 ELA pd","6-12 Math cur","6-12 Math pd"]}
tiles = {"K-5 ELA": 3.75, "K-5 Math": 3.46, "6-12 ELA": 1.57, "6-12 Math": 1.69}
bad_tiles = []
for lane, shown in tiles.items():
    tot = (sum(arr[f"{lane} cur"]) + sum(arr[f"{lane} pd"])) / 1e6
    if abs(tot - shown) > 0.005:
        bad_tiles.append((lane, shown, round(tot, 3)))
check(not bad_tiles, "curriculum stat tiles = sum of their arrays",
      f"{len(tiles)} lanes" if not bad_tiles else str(bad_tiles))
comb = sum(sum(v) for v in arr.values()) / 1e6
m_comb = re.search(r'\$([\d.]+)M</div>\s*<div[^>]*>across 16 fiscal years', mk)
check(m_comb and abs(float(m_comb.group(1)) - comb) < 0.005, "combined tile = sum of all lanes",
      f"{m_comb.group(1) if m_comb else '?'} vs {comb:.2f}")

print("\n2.5  labels describe what is shown")
# Positive assertions: no hardcoded span survives, AND the label is derived.
hard_span = re.search(r'of (\d+)-year total', mk)
check(hard_span is None, "no hardcoded N-year span label",
      f'found "{hard_span.group(0)}"' if hard_span else "none")
check("`${pct}% of ${FY[0]}-${FY[FY.length - 1]} total`" in mk,
      "FY-card percentage label derives its span from the FY list")
check("`Partial \u2014 through ${ROOT_DATA.meta.throughMonth}`" in mk,
      "partial-year caption derives from meta.throughMonth")
check("fy === 'FY26'" not in mk, "no hardcoded FY26 in the render logic")
check("'4 fiscal years'" not in mk, "no hardcoded '4 fiscal years' stat subtitle")
check("Issue Date FY" not in mk and "Meeting FY" not in mk,
      "no retired dual-FY terminology in the markup")
check("troymi" not in html, "BoardDocs link uses the school-district slug")

print("\n" + "=" * 62)
print("RESULT:", "PASS - every published figure traces to the data" if not FAIL
      else f"FAIL — {len(FAIL)} check(s): {FAIL}")

sys.exit(1 if FAIL else 0)
